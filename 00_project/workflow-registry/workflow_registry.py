#!/usr/bin/env python3
"""741 workflow registry: which workflow ran for which client, state, approval, last run, links.
Source of truth: registry.json (this folder) plus 00_project/cro-registry/registry.json (CRO pages).
Output: ~/Desktop/Workflow-Uebersicht.xlsx (plain-text links, no hyperlink objects, opens in Numbers and Excel).

Commands (run with the framework Python, it has openpyxl):
  workflow_registry.py build
  workflow_registry.py list
  workflow_registry.py add-client --name "Name" [--type Kunde|Lead|Partner] [--drive URL] [--local PATH]
  workflow_registry.py upsert --client "Name" --workflow <id> --status "<status>" [--mode Echt|Trockenlauf]
                              [--date YYYY-MM-DD] [--local PATH] [--drive URL] [--result PATH] [--note "text"]
  workflow_registry.py approve --client "Name" --workflow <id> [--date YYYY-MM-DD] [--note "text"]
  workflow_registry.py approve --cro-id <id from cro-registry> [--date YYYY-MM-DD]
  workflow_registry.py sync [--quiet]     automatic: reads run files the workflows write, updates the registry, rebuilds the Excel (Stop hook runs it after every turn)
  workflow_registry.py add-workflow --workflow <id> --name "Label"
Workflow ids: cro, cro_b, unburning, phone, audit, pool, citations, backlink
Status words: in Arbeit | Trockenlauf | fertig, Freigabe offen | freigegeben | Archiv  (no run = not started)
`approve` is the automatic entry when Martin approves a workflow result."""
import json, os, sys, datetime, argparse, re, fcntl, contextlib, unicodedata
HERE = os.path.dirname(os.path.abspath(__file__))
REG = os.path.join(HERE, 'registry.json')
CRO = os.path.join(HERE, '..', 'cro-registry', 'registry.json')
ROOT = '/Users/martindrendel/741website/'
OUT = os.path.expanduser('~/Desktop/Workflow-Uebersicht.xlsx')
STATUS = ['in Arbeit', 'Trockenlauf', 'fertig, Freigabe offen', 'freigegeben', 'Archiv']
RANK = {'freigegeben': 5, 'fertig, Freigabe offen': 4, 'in Arbeit': 3, 'Trockenlauf': 2, 'Archiv': 1}
today = lambda: datetime.date.today().isoformat()
def load(p): return json.load(open(p, encoding='utf-8'))
def save(p, d):
    tmp = p + '.tmp%d' % os.getpid()
    with open(tmp, 'w', encoding='utf-8') as f: json.dump(d, f, ensure_ascii=False, indent=1)
    os.replace(tmp, p)
LOCKF = os.path.join(HERE, '.registry.lock')
@contextlib.contextmanager
def locked():
    """One writer at a time. Several chats can end a turn at the same moment."""
    f = open(LOCKF, 'w'); fcntl.flock(f, fcntl.LOCK_EX)
    try: yield
    finally: fcntl.flock(f, fcntl.LOCK_UN); f.close()
def furl(p):
    if not p: return ''
    if p.startswith('http'): return p
    return 'file://' + (p if p.startswith('/') else ROOT + p)
def norm(s): return (s or '').strip().lower()

def cro_runs(reg):
    """CRO pages from cro-registry become runs. Client name is matched to a client in registry.json."""
    out = []
    if not os.path.exists(CRO): return out
    for e in load(CRO)['entries']:
        st = e.get('status', '')
        status = 'Archiv' if 'Archiv' in st else ('freigegeben' if e.get('approved') else 'fertig, Freigabe offen')
        label = e['page'] + ((' ' + e['version'].split()[0]) if e.get('version') else '') + (' (Archiv)' if status == 'Archiv' else '')
        out.append({'client': e['client'], 'workflow': 'cro', 'page': label, 'cro_id': e['id'], 'status': status,
                    'approved': bool(e.get('approved')), 'approved_date': e.get('approved_date', ''), 'mode': 'Echt',
                    'last_run': e.get('date', ''), 'drive': e.get('drive', ''), 'local': e.get('folder', ''),
                    'result': e.get('prototype', ''), 'score': e.get('score', ''), 'note': e.get('note', '')})
    return out

def all_runs(reg): return reg['runs'] + cro_runs(reg)

def cell_text(runs):
    if not runs: return '–'
    best = max(runs, key=lambda r: (RANK.get(r['status'], 0), r.get('last_run', '')))
    status_map = {'freigegeben': 'FREIGEGEBEN', 'fertig, Freigabe offen': 'fertig, Freigabe offen', 'in Arbeit': 'in Arbeit',
                  'Trockenlauf': 'Trockenlauf', 'Archiv': 'Archiv', 'drafted': 'in Arbeit'}
    t = status_map.get(best['status'], best['status'])
    d = best.get('approved_date') if best['status'] == 'freigegeben' and best.get('approved_date') else best.get('last_run', '')
    n = len([r for r in runs if r['status'] != 'Archiv'])
    return t + (' · ' + d if d else '') + (f' · {n} Seiten' if n > 1 else '')

def build():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    reg = load(REG); runs = all_runs(reg)
    wf = {w['id']: w for w in reg['workflows']}
    order = [w['id'] for w in reg['workflows'] if w.get('in_overview', True)]
    wb = Workbook()
    hdr = PatternFill('solid', fgColor='1F2A44'); hf = Font(bold=True, color='FFFFFF')
    fills = {'freigegeben': 'C6EFCE', 'fertig, Freigabe offen': 'FFEB9C', 'in Arbeit': 'FCE4D6', 'Trockenlauf': 'DDEBF7', 'Archiv': 'EDEDED', None: 'FFFFFF'}
    thin = Side(style='thin', color='D9D9D9'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    # --- sheet 1 overview
    ws = wb.active; ws.title = 'Übersicht'
    ws.append(['Kunde', 'Typ'] + [wf[i]['label'] for i in order])
    for c in ws[1]: c.fill = hdr; c.font = hf; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = box
    ws.append(['So starten', ''] + [wf[i].get('short', '') for i in order])
    for c in ws[2]: c.fill = PatternFill('solid', fgColor='E8ECF5'); c.font = Font(italic=True, size=10); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = box
    for cl in reg['clients']:
        row = [cl['name'], cl.get('type', '')]
        sts = []
        for i in order:
            rs = [r for r in runs if norm(r['client']) == norm(cl['name']) and r['workflow'] == i]
            row.append(cell_text(rs)); sts.append(max(rs, key=lambda r: RANK.get(r['status'], 0))['status'] if rs else None)
        ws.append(row)
        for j, s in enumerate(sts):
            c = ws.cell(row=ws.max_row, column=3 + j); c.fill = PatternFill('solid', fgColor=fills.get(s, 'FFFFFF')); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = box
        for j in (1, 2): ws.cell(row=ws.max_row, column=j).alignment = Alignment(vertical='top'); ws.cell(row=ws.max_row, column=j).border = box
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 9
    from openpyxl.utils import get_column_letter
    for k in range(len(order)): ws.column_dimensions[get_column_letter(3 + k)].width = 26
    ws.row_dimensions[1].height = 34; ws.row_dimensions[2].height = 62; ws.freeze_panes = 'C3'
    ws.append([]); ws.append(['Legende: grün = von Martin freigegeben · gelb = fertig, Freigabe offen · orange = in Arbeit · blau = Trockenlauf (Test) · grau = Archiv · – = nicht durchgeführt'])
    ws.append(['Neue Freigaben trägt Claude automatisch ein (Befehl `approve`). Alle Befehle: Blatt Befehle. Stand: ' + today()])
    # --- sheet 2 links
    wl = wb.create_sheet('Links')
    wl.append(['Kunde', 'Workflow', 'Seite oder Lauf', 'Status', 'Freigabe von Martin', 'Letzter Lauf', 'Ordner Drive', 'Ordner lokal', 'Ergebnis (Datei oder Link)', 'Notiz'])
    for c in wl[1]: c.fill = hdr; c.font = hf; c.border = box
    for cl in reg['clients']:
        rs = [r for r in runs if norm(r['client']) == norm(cl['name'])]
        first = True
        # client folder row
        wl.append([cl['name'], 'Kundenordner', '', '', '', '', cl.get('drive', ''), furl(cl.get('local', '')), '', cl.get('note', '')])
        for c in wl[wl.max_row]: c.font = Font(bold=True)
        for r in sorted(rs, key=lambda r: (order.index(r['workflow']) if r['workflow'] in order else 99, r.get('last_run', ''))):
            ap = ('ja, ' + r['approved_date']) if r.get('approved') and r.get('approved_date') else ('ja' if r.get('approved') else 'offen')
            wl.append(['', wf[r['workflow']]['label'], (r.get('page', '') + (' · ' if r.get('page') and r.get('mode') == 'Trockenlauf' else '') + ('Trockenlauf' if r.get('mode') == 'Trockenlauf' else '')).strip(),
                       r['status'], ap, r.get('last_run', ''), r.get('drive', ''), furl(r.get('local', '')), furl(r.get('result', '')), r.get('note', '')])
            wl.cell(row=wl.max_row, column=4).fill = PatternFill('solid', fgColor=fills.get(r['status'], 'FCE4D6'))
        wl.append([])
    for col, w in zip('ABCDEFGHIJ', [22, 30, 28, 22, 18, 12, 62, 70, 80, 60]): wl.column_dimensions[col].width = w
    wl.freeze_panes = 'C2'
    # --- sheet 3 workflows
    ww = wb.create_sheet('Workflows')
    ww.append(['Workflow', 'Auch genannt', 'Reife des Workflows', 'Workflow freigegeben (Martin)', 'Slash-Befehl zum Starten', 'Stand des Slash-Befehls', 'Satz im Chat (geht heute)', 'Weiterarbeiten', 'Ergebnis ansehen', 'Installieren / Terminal', 'Dokument (Pfad)', 'Skill (Pfad)', 'Läufe', 'davon freigegeben', 'Notiz'])
    for c in ww[1]: c.fill = hdr; c.font = hf; c.border = box
    for w in reg['workflows']:
        rs = [r for r in runs if r['workflow'] == w['id']]
        ww.append([w['label'], w.get('alias', ''), w.get('maturity', ''), w.get('approved', ''), w.get('slash', ''), w.get('slash_status', ''), w.get('chat', ''), w.get('resume', ''), w.get('view', ''), w.get('install', ''), furl(w.get('doc', '')), furl(w.get('skill', '')), len(rs), len([r for r in rs if r['status'] == 'freigegeben']), w.get('note', '')])
    for col, w in zip('ABCDEFGHIJKLMNO', [30, 30, 32, 24, 40, 30, 46, 34, 52, 58, 62, 62, 8, 12, 60]): ww.column_dimensions[col].width = w
    ww.freeze_panes = 'B2'
    wc = wb.create_sheet('Befehle')
    wc.append(['Was du willst', 'Aktion', 'Befehl oder Satz', 'Wo', 'Hinweis'])
    for c in wc[1]: c.fill = hdr; c.font = hf; c.border = box
    for row in reg.get('commands', []):
        wc.append([row[0], row[1], row[2], row[3], row[4]])
        wc.cell(row=wc.max_row, column=1).font = Font(bold=True)
    for col, w in zip('ABCDE', [14, 44, 80, 26, 90]): wc.column_dimensions[col].width = w
    wc.freeze_panes = 'A2'
    for sh in (ws, wl, ww, wc):
        for row in sh.iter_rows(min_row=2):
            for c in row:
                if c.alignment is None or not c.alignment.wrap_text: c.alignment = Alignment(wrap_text=True, vertical='top')
    tmpx = OUT[:-5] + '.tmp.xlsx'; wb.save(tmpx); os.replace(tmpx, OUT)
    print(f'geschrieben: {OUT} · {len(reg["clients"])} Kunden · {len(runs)} Läufe · {len(reg["workflows"])} Workflows')

def find_client(reg, name):
    for c in reg['clients']:
        if norm(c['name']) == norm(name) or norm(name) in norm(c['name']): return c
    return None

def cmd_upsert(a, approve=False):
    reg = load(REG)
    if a.cro_id:
        d = load(CRO)
        for e in d['entries']:
            if e['id'] == a.cro_id:
                e['approved'] = True; e['approved_date'] = a.date or today(); e['status'] = 'Freigegeben'
                save(CRO, d); build(); print('CRO approved:', a.cro_id); return
        sys.exit('cro-id not found')
    cl = find_client(reg, a.client)
    if not cl: sys.exit('client not found, use add-client first: ' + str(a.client))
    if a.workflow not in {w['id'] for w in reg['workflows']}: sys.exit('unknown workflow id')
    run = next((r for r in reg['runs'] if norm(r['client']) == norm(cl['name']) and r['workflow'] == a.workflow), None)
    if not run:
        run = {'client': cl['name'], 'workflow': a.workflow, 'status': 'in Arbeit', 'approved': False, 'approved_date': '', 'mode': 'Echt', 'last_run': today(), 'drive': '', 'local': '', 'result': '', 'note': ''}
        reg['runs'].append(run)
    if approve:
        run['status'] = 'freigegeben'; run['approved'] = True; run['approved_date'] = a.date or today()
    else:
        if a.status:
            if a.status not in STATUS: sys.exit('status must be one of ' + ', '.join(STATUS))
            run['status'] = a.status
            if a.status != 'freigegeben': run['approved'] = False; run['approved_date'] = ''
    for k in ('mode', 'drive', 'local', 'result', 'note'):
        v = getattr(a, k, None)
        if v: run[k] = v
    run['last_run'] = a.date or run.get('last_run') or today()
    save(REG, reg); build(); print('ok:', cl['name'], a.workflow, run['status'])

# ---------------------------------------------------------------- automatic sync (2026-10-11)
ASSETS = os.path.join(ROOT, '00_project', 'assets')
SCAN_ROOTS = [ASSETS, os.path.join(ROOT, '00_project', 'prospects')]
SKIP_DIRS = {'node_modules', '.git', 'evidence', 'scripts', 'images', '_archiv', '_drive', 'screenshots', 'falcon'}
MARKERS = {'_run.json', 'run.json', 'matrix.json', '_status.md', 'harm-register.json'}

def fold(s):
    s = (s or '').lower()
    for a, b in (('ä', 'ae'), ('ö', 'oe'), ('ü', 'ue'), ('ß', 'ss')): s = s.replace(a, b)
    s = ''.join(c for c in unicodedata.normalize('NFKD', s) if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()

def rel(p): return os.path.relpath(p, ROOT)
def iso_mtime(paths):
    t = [os.path.getmtime(p) for p in paths if os.path.exists(p)]
    return datetime.date.fromtimestamp(max(t)).isoformat() if t else today()
def status_dates(p):
    ds = [d for d in re.findall(r'^(\d{4}-\d{2}-\d{2})', open(p, encoding='utf-8', errors='ignore').read(), re.M) if d <= today()]
    return max(ds) if ds else iso_mtime([p])
def jload(p):
    try: return json.load(open(p, encoding='utf-8'))
    except Exception: return {}

def collect_markers():
    found = {}
    for r, ds, fs in (x for root in SCAN_ROOTS for x in os.walk(root)):
        ds[:] = [d for d in ds if d not in SKIP_DIRS]
        for f in fs:
            if f in MARKERS: found.setdefault(f, []).append(os.path.join(r, f))
    return found

def slug_of(path):
    """00_project/assets/<slug>/... -> slug (also un-burning/<slug>/, audit/<slug>/)"""
    parts = rel(path).split(os.sep)[2:]
    if parts and parts[0] in ('un-burning', 'audit'): return parts[1] if len(parts) > 1 else ''
    return parts[0] if parts else ''

def client_for(reg, name=None, slug=None, hint=''):
    if name:
        n = fold(name)
        for c in reg['clients']:
            cn = fold(c['name'])
            if n == cn or n in cn or cn in n: return c, False
    if slug:
        s = fold(slug).replace(' ', '')
        for c in reg['clients']:
            toks = fold(c['name']).split() + fold(c.get('local', '')).split()
            if s in [t for t in toks] or s in fold(c['name']).replace(' ', ''): return c, False
    nm = name or (slug or 'unbekannt').replace('-', ' ').title()
    c = {'name': nm, 'type': 'Lead', 'drive': '', 'local': '', 'note': 'automatisch angelegt (Ordner ' + hint + '), bitte prüfen'}
    reg['clients'].append(c)
    return c, True

def detect(reg):
    """Return a list of detected runs. Each detector reads files that the workflows write anyway."""
    m = collect_markers(); out = []
    # explicit marker, any workflow: _run.json  (schema in README.md)
    for p in m.get('_run.json', []):
        d = jload(p)
        if d.get('workflow'):
            out.append({'explicit': True, 'client_name': d.get('client'), 'slug': slug_of(p), 'workflow': d['workflow'],
                        'status': d.get('status') or 'in Arbeit', 'mode': d.get('mode') or 'Echt',
                        'last_run': d.get('last_run') or iso_mtime([p]), 'local': d.get('local') or rel(os.path.dirname(p)),
                        'result': d.get('result', ''), 'drive': d.get('drive', ''), 'note': d.get('note', ''), 'page': d.get('page', ''), 'hint': rel(p)})
    for p in m.get('run.json', []):
        folder = os.path.dirname(p); d = jload(p); rp = rel(p)
        if '/audit/' in rp.replace(os.sep, '/') and d.get('subject_name'):
            rep = os.path.join(folder, 'report')
            done = d.get('status') in ('done', 'complete', 'completed', 'final', 'reported') or (os.path.isdir(rep) and any(f.endswith(('.html', '.xlsx')) for f in os.listdir(rep)))
            res = next((rel(os.path.join(rep, f)) for f in ('internal.html', 'befunde.xlsx') if os.path.exists(os.path.join(rep, f))), rel(os.path.join(folder, 'findings.json')))
            bits = [f"Lauf {d.get('run_id')}"]
            if d.get('score_total') is not None: bits.append(f"Score {d['score_total']}/100")
            if d.get('verdict'): bits.append(f"Verdikt {d['verdict']}")
            out.append({'client_name': d['subject_name'], 'slug': d.get('subject_slug'), 'workflow': 'audit',
                        'status': 'fertig, Freigabe offen' if done else 'in Arbeit', 'mode': 'Trockenlauf' if 'dry' in str(d.get('mode')).lower() else 'Echt',
                        'last_run': iso_mtime([p] + [os.path.join(rep, f) for f in (os.listdir(rep) if os.path.isdir(rep) else [])]),
                        'local': rel(folder), 'result': res, 'note': ', '.join(bits) + ' (automatisch erkannt)', 'hint': rp})
        elif '/un-burning/' in rp.replace(os.sep, '/') and d.get('client'):
            dry = 'dry' in str(d.get('mode')).lower()
            base = os.path.join(os.path.dirname(folder), 'baseline', 'BASELINE-REPORT.md')
            res = rel(base) if os.path.exists(base) else rel(os.path.join(folder, 'harm-register.json'))
            out.append({'client_name': d['client'], 'slug': slug_of(p), 'workflow': 'unburning',
                        'status': 'Trockenlauf' if dry else ('fertig, Freigabe offen' if d.get('finished_at') else 'in Arbeit'),
                        'mode': 'Trockenlauf' if dry else 'Echt', 'last_run': iso_mtime([p, os.path.join(folder, 'harm-register.json')]),
                        'local': rel(folder), 'result': res, 'note': f"Lauf {d.get('run_id')} (automatisch erkannt)", 'hint': rp})
    for p in m.get('matrix.json', []):
        d = jload(p); folder = os.path.dirname(p)
        if isinstance(d, dict) and d.get('client'):
            dry = 'dryrun' in rel(p).lower() or 'dry-run' in rel(p).lower()
            ev = os.path.join(folder, 'evidence-index.md')
            out.append({'client_name': d['client'], 'slug': slug_of(p), 'workflow': 'phone', 'status': 'Trockenlauf' if dry else 'in Arbeit',
                        'mode': 'Trockenlauf' if dry else 'Echt', 'last_run': d.get('as_of') or iso_mtime([p]), 'local': rel(folder),
                        'result': rel(ev) if os.path.exists(ev) else rel(p), 'note': 'Maßnahmen-Matrix, Katalog ' + str(d.get('catalog_version', '?')) + ' (automatisch erkannt)', 'hint': rel(p)})
    for p in m.get('_status.md', []):
        folder = os.path.dirname(p); name = os.path.basename(folder); parent = os.path.basename(os.path.dirname(folder))
        txt = open(p, encoding='utf-8', errors='ignore').read(); last = status_dates(p)
        if name == 'content-pool':
            rep = os.path.join(folder, '05-report', 'content-pool-report.html')
            up = os.path.exists(os.path.join(folder, '_drive', '_status.md'))
            out.append({'client_name': None, 'slug': slug_of(p), 'workflow': 'pool', 'status': 'fertig, Freigabe offen' if (os.path.exists(rep) and up) else 'in Arbeit',
                        'mode': None, 'last_run': last, 'local': rel(folder), 'result': rel(rep) if os.path.exists(rep) else '', 'note': 'Pool aus _status.md erkannt', 'hint': rel(p)})
        elif name == 'cro-b':
            out.append({'client_name': None, 'slug': slug_of(p), 'workflow': 'cro_b', 'status': 'fertig, Freigabe offen' if 'handoff' in txt.lower() else 'in Arbeit',
                        'mode': None, 'last_run': last, 'local': rel(folder), 'result': '', 'note': 'Track B aus _status.md erkannt', 'hint': rel(p)})
        elif name == 'cro' or parent == 'cro':
            proto = os.path.join(folder, '07-prototype', 'index.html')
            out.append({'client_name': None, 'slug': slug_of(p), 'workflow': 'cro', 'status': 'fertig, Freigabe offen' if 'step 7 done' in txt.lower() else 'in Arbeit',
                        'mode': None, 'last_run': last, 'local': rel(folder), 'page': '' if name == 'cro' else name,
                        'result': rel(proto) if os.path.exists(proto) else rel(os.path.join(folder, '04-copy', 'copy.md')), 'note': 'CRO-Seite aus _status.md erkannt', 'hint': rel(p), 'cro': True})
    return out

def cmd_sync(quiet=False):
    with locked():
        reg = load(REG); changes = []
        covered = {os.path.normpath(e.get('folder', '')) for e in (load(CRO)['entries'] if os.path.exists(CRO) else [])}
        wfids = {w['id'] for w in reg['workflows']}
        best = {}
        for d in detect(reg):
            c, new = client_for(reg, d.get('client_name'), d.get('slug'), d.get('hint', ''))
            if new: changes.append('neuer Kunde: ' + c['name'])
            if d['workflow'] not in wfids:
                reg['workflows'].append({'id': d['workflow'], 'label': d['workflow'], 'alias': '', 'maturity': 'neu, automatisch angelegt, bitte Eintrag ergänzen', 'approved': 'offen', 'doc': '', 'skill': '', 'note': '', 'slash': '', 'slash_status': '', 'chat': '', 'resume': '', 'view': '', 'install': '', 'short': ''})
                wfids.add(d['workflow']); changes.append('neuer Workflow: ' + d['workflow'])
            if d.get('cro') and os.path.normpath(d['local']) in covered:
                continue  # page lives in cro-registry (manual entry), nothing to add
            k = (norm(c['name']), d['workflow'], d.get('page', '') if d.get('cro') else '')
            if k not in best or d['last_run'] >= best[k][1]['last_run']: best[k] = (c, d)
        for (cn, wf, pg), (c, d) in best.items():
            run = next((r for r in reg['runs'] if norm(r['client']) == cn and r['workflow'] == wf and (r.get('page', '') == pg or not d.get('cro'))), None)
            tag = f"{c['name']} / {wf}"
            if not run:
                reg['runs'].append({'client': c['name'], 'workflow': wf, 'status': d['status'], 'approved': False, 'approved_date': '', 'mode': d.get('mode') or 'Echt',
                                    'last_run': d['last_run'], 'drive': d.get('drive', ''), 'local': d.get('local', ''), 'result': d.get('result', ''), 'note': d.get('note', ''),
                                    **({'page': d['page']} if d.get('page') else {})})
                changes.append('neuer Lauf: ' + tag + ' · ' + d['status']); continue
            if d.get('explicit'):
                for k2 in ('status', 'mode', 'local', 'result', 'drive', 'note'):
                    if d.get(k2) and run.get(k2) != d[k2] and not (k2 == 'status' and run.get('approved')): run[k2] = d[k2]; changes.append(f'{tag}: {k2} aus _run.json')
            else:
                locked_status = run.get('approved') or run['status'] in ('freigegeben', 'Archiv')
                can_raise = (run.get('mode') != 'Trockenlauf' or d.get('mode') == 'Echt')
                if not locked_status and can_raise and RANK.get(d['status'], 0) > RANK.get(run['status'], 0):
                    old = run['status']; run['status'] = d['status']
                    if d.get('mode') == 'Echt': run['mode'] = 'Echt'
                    if d.get('note'): run['note'] = d['note'] + ((' · vorher: ' + run.get('note', '')[:90]) if run.get('note') else '')
                    changes.append(f'{tag}: {old} -> {d["status"]}')
                for k2 in ('local', 'result'):
                    if d.get(k2) and not run.get(k2): run[k2] = d[k2]; changes.append(f'{tag}: {k2} ergänzt')
            if d['last_run'] > (run.get('last_run') or ''): run['last_run'] = d['last_run']; changes.append(f'{tag}: letzter Lauf {d["last_run"]}')
        stale_x = not os.path.exists(OUT) or os.path.getmtime(OUT) < max(os.path.getmtime(REG), os.path.getmtime(CRO) if os.path.exists(CRO) else 0)
        if changes: save(REG, reg)
        if changes or stale_x: build_quiet = quiet; build()
        if not quiet:
            print('\n'.join(changes) if changes else 'sync: keine Änderung')
            if os.path.exists(os.path.expanduser('~/Desktop/~$Workflow-Uebersicht.xlsx')): print('Hinweis: Die Excel ist offen. Schließen und neu öffnen, damit du die Änderungen siehst.')
    return changes

def cmd_add_workflow(a):
    with locked():
        reg = load(REG)
        if any(w['id'] == a.workflow for w in reg['workflows']): sys.exit('exists')
        reg['workflows'].append({'id': a.workflow, 'label': a.name or a.workflow, 'alias': '', 'maturity': a.note or 'neu', 'approved': 'offen', 'doc': a.result or '', 'skill': '', 'note': '', 'slash': '', 'slash_status': '', 'chat': '', 'resume': '', 'view': '', 'install': '', 'short': ''})
        save(REG, reg); build()

if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('cmd'); p.add_argument('--quiet', action='store_true')
    for k in ('name', 'type', 'drive', 'local', 'client', 'workflow', 'status', 'mode', 'date', 'result', 'note', 'cro_id'):
        p.add_argument('--' + k.replace('_', '-'), dest=k)
    a = p.parse_args()
    if a.cmd == 'build': build()
    elif a.cmd == 'list':
        for r in all_runs(load(REG)): print(r['client'], '|', r['workflow'], '|', r['status'], '|', r.get('last_run'))
    elif a.cmd == 'add-client':
        reg = load(REG)
        if find_client(reg, a.name): sys.exit('exists')
        reg['clients'].append({'name': a.name, 'type': a.type or 'Kunde', 'drive': a.drive or '', 'local': a.local or '', 'note': a.note or ''}); save(REG, reg); build()
    elif a.cmd == 'upsert':
        with locked(): cmd_upsert(a)
    elif a.cmd == 'approve':
        with locked(): cmd_upsert(a, approve=True)
    elif a.cmd == 'sync':
        try: cmd_sync(quiet=a.quiet)
        except Exception as e:
            if not a.quiet: raise
            sys.exit(0)  # a hook never fails a turn
    elif a.cmd == 'add-workflow': cmd_add_workflow(a)
    else: sys.exit('unknown command')
